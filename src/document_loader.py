# src/document_loader.py
import docx                     # reads .docx files
from PIL import Image           # reads .png image files
import os


def load_document(filepath: str):
    """
    Given a file path, return:
      - (str, "text")  for DOCX files  → returns extracted plain text
      - (Image, "image") for PNG files → returns PIL Image object
    """
    ext = os.path.splitext(filepath)[1].lower()   # get extension like .docx or .png
    if ext == ".docx":
        doc = docx.Document(filepath)
        lines = []

        # Read all paragraph texts
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text)

        # Also read table cells
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        lines.append(cell.text.strip())

        text = "\n".join(lines)
        return text, "text"

    elif ext == ".png":
        image = Image.open(filepath)
        return image, "image"
    else:
        raise ValueError(f"Unsupported file type: {ext}")